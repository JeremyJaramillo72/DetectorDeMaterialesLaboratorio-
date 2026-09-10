# -*- coding: utf-8 -*-
"""
Generador de Manuales Individuales por Equipo en Formato Tabla y Paleta Blanco/Negro
Laboratorio de Bromatología - Universidad Técnica Estatal de Quevedo (UTEQ)
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Importar lista de equipos de generate_manual_docx
sys.path.insert(0, os.path.dirname(__file__))
from generate_manual_docx import EQUIPOS, IMAGES_DIR, WORKSPACE_DIR

OUTPUT_DIR = os.path.join(WORKSPACE_DIR, "knowledge_base", "manuales_individuales")
os.makedirs(OUTPUT_DIR, exist_ok=True)

# Paleta estricta Blanco y Negro / Escala de Grises
COLOR_BLACK = RGBColor(0x00, 0x00, 0x00)
COLOR_DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
COLOR_GRAY = RGBColor(0x66, 0x66, 0x66)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEX_HEADER_BG = "000000"       # Negro puro para cabeceras
HEX_SUBHEADER_BG = "262626"    # Gris muy oscuro para subtítulos de tabla
HEX_ZEBRA_BG = "F4F4F4"        # Gris muy tenue para filas alternas
HEX_WHITE_BG = "FFFFFF"        # Blanco puro
HEX_BORDER = "333333"          # Borde gris oscuro definido

def set_cell_background(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tc_pr.append(tcMar)

def format_monochrome_table(table, col_widths, col_alignments=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    for i, row in enumerate(table.rows):
        is_header = (i == 0)
        for j, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[j])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=90, bottom=90, left=130, right=130)
            
            # Alineación por columna si se especifica
            if col_alignments and j < len(col_alignments):
                for p in cell.paragraphs:
                    p.alignment = col_alignments[j]
                    
            if is_header:
                set_cell_background(cell, HEX_HEADER_BG)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for r in p.runs:
                        r.bold = True
                        r.font.name = "Calibri"
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = COLOR_WHITE
            else:
                bg = HEX_ZEBRA_BG if i % 2 == 1 else HEX_WHITE_BG
                set_cell_background(cell, bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.name = "Calibri"
                        r.font.size = Pt(9.0)
                        r.font.color.rgb = COLOR_BLACK
                        
            # Bordes negro/gris oscuro
            tc_pr = cell._tc.get_or_add_tcPr()
            b = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                          f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
                          f'<w:left w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
                          f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
                          f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
                          f'</w:tcBorders>')
            tc_pr.append(b)

def add_monochrome_callout(doc, title, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "EAEAEA")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="none"/>'
                        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="000000"/>'
                        f'<w:bottom w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'</w:tcBorders>')
    tc_pr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"[!] {title}: ")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(9.5)
    run_t.font.color.rgb = COLOR_BLACK
    
    run_c = p.add_run(text)
    run_c.font.name = "Calibri"
    run_c.font.size = Pt(9.0)
    run_c.font.color.rgb = COLOR_DARK_GRAY
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

def add_section_title(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.color.rgb = COLOR_BLACK

FILE_NAMES = [
    "Manual_01_Destilador_Arrastre_Vapor_Kjeldahl.docx",
    "Manual_02_Analizador_Fibra_Cruda_DosiFiber.docx",
    "Manual_03_Viscosimetro_Brookfield_DVE.docx",
    "Manual_04_Microscopio_Trinocular_Zeiss.docx",
    "Manual_05_Destilador_Nitrogeno_Proteinas_Fisher.docx",
    "Manual_06_Destilador_Agua_Continuo_Metalico.docx",
    "Manual_07_Sistema_Tratamiento_Desionizacion_Agua.docx",
    "Manual_08_Bomba_Vacio_Recirculacion_Agua.docx"
]

def generate_individual_documents():
    print(f"Generating 8 individual documents in monochrome palette...")
    
    for idx, eq in enumerate(EQUIPOS):
        doc = docx.Document()
        
        # Márgenes
        for sec in doc.sections:
            sec.top_margin = Inches(0.8)
            sec.bottom_margin = Inches(0.8)
            sec.left_margin = Inches(1.0)
            sec.right_margin = Inches(1.0)
            sec.page_width = Inches(8.5)
            sec.page_height = Inches(11.0)
            
        filename = FILE_NAMES[idx]
        filepath = os.path.join(OUTPUT_DIR, filename)
        
        # ---------------------------------------------------------
        # ENCABEZADO INSTITUCIONAL
        # ---------------------------------------------------------
        p_head = doc.add_paragraph()
        p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_head.paragraph_format.space_before = Pt(4)
        p_head.paragraph_format.space_after = Pt(2)
        r_u = p_head.add_run("UNIVERSIDAD TÉCNICA ESTATAL DE QUEVEDO\n")
        r_u.bold = True
        r_u.font.name = "Calibri"
        r_u.font.size = Pt(13)
        r_u.font.color.rgb = COLOR_BLACK
        
        r_f = p_head.add_run("FACULTAD DE CIENCIAS PECUARIAS Y BIOLÓGICAS — LABORATORIO DE BROMATOLOGÍA\n")
        r_f.font.name = "Calibri"
        r_f.font.size = Pt(9.5)
        r_f.font.color.rgb = COLOR_DARK_GRAY
        
        r_p = p_head.add_run("MANUAL TÉCNICO OPERATIVO Y BASE DE CONOCIMIENTO RAG (YOLO11)")
        r_p.bold = True
        r_p.font.name = "Calibri"
        r_p.font.size = Pt(9.0)
        r_p.font.color.rgb = COLOR_GRAY
        
        # Línea horizontal divisoria
        p_div = doc.add_paragraph()
        p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_div.paragraph_format.space_before = Pt(2)
        p_div.paragraph_format.space_after = Pt(8)
        r_div = p_div.add_run("―" * 55)
        r_div.font.color.rgb = COLOR_BLACK
        
        # TÍTULO DEL EQUIPO
        p_tit = doc.add_paragraph()
        p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_tit.paragraph_format.space_before = Pt(4)
        p_tit.paragraph_format.space_after = Pt(2)
        r_tit = p_tit.add_run(f"EQUIPO N° {eq['num']}: {eq['nombre_oficial'].upper()}")
        r_tit.bold = True
        r_tit.font.name = "Calibri"
        r_tit.font.size = Pt(14)
        r_tit.font.color.rgb = COLOR_BLACK
        
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(8)
        r_sub = p_sub.add_run(f"Clase YOLO11: [{eq['clase_yolo']}]  |  Modelo: {eq['modelo']}")
        r_sub.font.name = "Calibri"
        r_sub.font.size = Pt(10)
        r_sub.font.color.rgb = COLOR_DARK_GRAY
        
        # FOTOGRAFÍA OFICIAL
        img_path = os.path.join(IMAGES_DIR, eq['imagen'])
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(4)
            p_img.paragraph_format.space_after = Pt(2)
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Inches(3.2))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(8)
            r_cap = p_cap.add_run(f"Fotografía 1.0: Vista frontal en el Laboratorio de Bromatología — {eq['nombre_oficial']}")
            r_cap.font.name = "Calibri"
            r_cap.font.size = Pt(8.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = COLOR_GRAY
            
        # ---------------------------------------------------------
        # TABLA 1: FICHA DE IDENTIFICACIÓN Y ESPECIFICACIONES TÉCNICAS
        # ---------------------------------------------------------
        add_section_title(doc, "1. Ficha de Identificación y Especificaciones Técnicas")
        
        t1_data = [
            ("Parámetro Instrumental", "Especificación Oficial Verificada"),
            ("Nombre Oficial en Laboratorio", eq['nombre_oficial']),
            ("Etiqueta Roboflow / YOLO11", eq['clase_yolo']),
            ("Fabricante y Procedencia", eq['fabricante']),
            ("Modelo y Referencia de Catálogo", eq['modelo']),
            ("Ubicación en Bromatología UTEQ", eq['ubicacion']),
            ("Alimentación Eléctrica", eq['alimentacion']),
            ("Requerimientos de Servicios", eq['servicios']),
            ("Precio Comercial Aproximado (USD)", eq.get('precio_aproximado', 'N/A')),
            ("Función Primaria Analítica", eq['principio'].split('\n\n')[0])
        ]
        
        tbl1 = doc.add_table(rows=len(t1_data), cols=2)
        for r_i, (c0_txt, c1_txt) in enumerate(t1_data):
            tbl1.cell(r_i, 0).paragraphs[0].text = c0_txt
            tbl1.cell(r_i, 1).paragraphs[0].text = c1_txt
            if r_i > 0:
                tbl1.cell(r_i, 0).paragraphs[0].runs[0].bold = True
        format_monochrome_table(tbl1, [2.3, 4.2])
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        
        # ---------------------------------------------------------
        # TABLA 2: FUNDAMENTO QUÍMICO-FÍSICO, REACCIONES Y CÁLCULOS
        # ---------------------------------------------------------
        add_section_title(doc, "2. Fundamento Analítico, Reacciones Químicas y Cálculos")
        
        t2_data = [("Mecanismo / Etapa", "Reacción Química, Principio Físico o Ecuación de Cálculo")]
        for r_tit, r_eq in eq['reacciones']:
            t2_data.append((r_tit, r_eq))
        for c_tit, c_eq in eq['calculos']:
            t2_data.append((f"Fórmula: {c_tit}", c_eq))
            
        tbl2 = doc.add_table(rows=len(t2_data), cols=2)
        for r_i, (c0_txt, c1_txt) in enumerate(t2_data):
            tbl2.cell(r_i, 0).paragraphs[0].text = c0_txt
            tbl2.cell(r_i, 1).paragraphs[0].text = c1_txt
            if r_i > 0:
                tbl2.cell(r_i, 0).paragraphs[0].runs[0].bold = True
                tbl2.cell(r_i, 1).paragraphs[0].runs[0].font.name = "Consolas"
                tbl2.cell(r_i, 1).paragraphs[0].runs[0].font.size = Pt(8.5)
        format_monochrome_table(tbl2, [2.2, 4.3])
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        
        # ---------------------------------------------------------
        # TABLA 3: ARQUITECTURA INSTRUMENTAL Y COMPONENTES CRÍTICOS
        # ---------------------------------------------------------
        add_section_title(doc, "3. Arquitectura Instrumental y Componentes Críticos")
        
        t3_data = [("N°", "Componente Instrumental", "Descripción y Función Técnica")]
        for comp_i, comp_txt in enumerate(eq['componentes']):
            if ":" in comp_txt:
                parts = comp_txt.split(":", 1)
                comp_name = parts[0].strip()
                comp_desc = parts[1].strip()
            else:
                words = comp_txt.split(" ", 3)
                comp_name = " ".join(words[:3])
                comp_desc = comp_txt
            t3_data.append((str(comp_i + 1), comp_name, comp_desc))
            
        tbl3 = doc.add_table(rows=len(t3_data), cols=3)
        for r_i, (c0_txt, c1_txt, c2_txt) in enumerate(t3_data):
            tbl3.cell(r_i, 0).paragraphs[0].text = c0_txt
            tbl3.cell(r_i, 1).paragraphs[0].text = c1_txt
            tbl3.cell(r_i, 2).paragraphs[0].text = c2_txt
            if r_i > 0:
                tbl3.cell(r_i, 1).paragraphs[0].runs[0].bold = True
        format_monochrome_table(tbl3, [0.4, 2.1, 4.0], [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        
        # ---------------------------------------------------------
        # TABLA 4: PROCEDIMIENTO OPERATIVO NORMALIZADO (SOP)
        # ---------------------------------------------------------
        add_section_title(doc, "4. Procedimiento Operativo Normalizado (PNO / SOP Paso a Paso)")
        
        t4_data = [("Fase Operativa", "Instrucciones de Ejecución Paso a Paso", "Puntos Críticos de Control")]
        puntos_criticos = [
            "Verificar agua de calderín y suministro de reactivos.",
            "Esperar estabilización térmica previa al ciclo.",
            "Comprobar hermeticidad y viraje de color / flujo.",
            "Desconexión oportuna y lectura volumétrica exacta.",
            "Descontaminación con agua desionizada y corte general."
        ]
        for f_i, (f_tit, f_desc) in enumerate(eq['sop']):
            crit = puntos_criticos[f_i] if f_i < len(puntos_criticos) else "Cumplir parámetros estandarizados."
            t4_data.append((f_tit, f_desc, crit))
            
        tbl4 = doc.add_table(rows=len(t4_data), cols=3)
        for r_i, (c0_txt, c1_txt, c2_txt) in enumerate(t4_data):
            tbl4.cell(r_i, 0).paragraphs[0].text = c0_txt
            tbl4.cell(r_i, 1).paragraphs[0].text = c1_txt
            tbl4.cell(r_i, 2).paragraphs[0].text = c2_txt
            if r_i > 0:
                tbl4.cell(r_i, 0).paragraphs[0].runs[0].bold = True
                tbl4.cell(r_i, 2).paragraphs[0].runs[0].font.size = Pt(8.5)
                tbl4.cell(r_i, 2).paragraphs[0].runs[0].font.italic = True
        format_monochrome_table(tbl4, [1.8, 3.4, 1.3])
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        
        # ---------------------------------------------------------
        # TABLA 5: BIOSEGURIDAD, EPP OBLIGATORIO Y MATRIZ DE RIESGOS
        # ---------------------------------------------------------
        add_section_title(doc, "5. Matriz de Bioseguridad, EPP Obligatorio y Prevención de Riesgos")
        
        t5_data = [("Elemento EPP", "Especificación Técnica Requerida", "Peligro Mitigado")]
        peligros_mitigados = [
            "Salpicaduras corrosivas al tronco y extremidades.",
            "Proyección ocular directa de reactivos químicos.",
            "Contacto dérmico con bases o ácidos concentrados.",
            "Quemaduras térmicas por contacto con piezas a > 80 °C.",
            "Resbalones, contacto con derrames químicos en suelo."
        ]
        for e_i, (e_nom, e_esp) in enumerate(eq['epp']):
            mit = peligros_mitigados[e_i] if e_i < len(peligros_mitigados) else "Exposición a riesgos de laboratorio."
            t5_data.append((e_nom, e_esp, mit))
            
        tbl5 = doc.add_table(rows=len(t5_data), cols=3)
        for r_i, (c0_txt, c1_txt, c2_txt) in enumerate(t5_data):
            tbl5.cell(r_i, 0).paragraphs[0].text = c0_txt
            tbl5.cell(r_i, 1).paragraphs[0].text = c1_txt
            tbl5.cell(r_i, 2).paragraphs[0].text = c2_txt
            if r_i > 0:
                tbl5.cell(r_i, 0).paragraphs[0].runs[0].bold = True
        format_monochrome_table(tbl5, [1.8, 3.2, 1.5])
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        
        # Alertas de seguridad en callout monocromático
        for rk in eq['riesgos']:
            add_monochrome_callout(doc, "ALERTA DE SEGURIDAD OPERATIVA", rk)
            
        # ---------------------------------------------------------
        # TABLA 6: GUÍAS DE PRÁCTICA UTEQ Y MÉTODOS NORMALIZADOS
        # ---------------------------------------------------------
        add_section_title(doc, "6. Guías de Práctica Oficiales UTEQ y Métodos Normalizados")
        
        t6_data = [("Ámbito / Aplicación", "Detalle de la Práctica de Laboratorio o Norma")]
        for g in eq['guias_uteq']:
            t6_data.append(("Práctica Curricular UTEQ", g))
        t6_data.append(("Normas Oficiales Aplicables", eq['normas']))
        
        tbl6 = doc.add_table(rows=len(t6_data), cols=2)
        for r_i, (c0_txt, c1_txt) in enumerate(t6_data):
            tbl6.cell(r_i, 0).paragraphs[0].text = c0_txt
            tbl6.cell(r_i, 1).paragraphs[0].text = c1_txt
            if r_i > 0:
                tbl6.cell(r_i, 0).paragraphs[0].runs[0].bold = True
        format_monochrome_table(tbl6, [2.0, 4.5])
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        
        # ---------------------------------------------------------
        # TABLA 7: PLAN DE MANTENIMIENTO Y CALIBRACIÓN
        # ---------------------------------------------------------
        add_section_title(doc, "7. Plan de Mantenimiento Preventivo y Calibración Metrológica")
        
        t7_data = [("Frecuencia", "Procedimiento de Mantenimiento / Calibración")]
        for m_frec, m_prot in eq['mantenimiento']:
            t7_data.append((m_frec, m_prot))
            
        tbl7 = doc.add_table(rows=len(t7_data), cols=2)
        for r_i, (c0_txt, c1_txt) in enumerate(t7_data):
            tbl7.cell(r_i, 0).paragraphs[0].text = c0_txt
            tbl7.cell(r_i, 1).paragraphs[0].text = c1_txt
            if r_i > 0:
                tbl7.cell(r_i, 0).paragraphs[0].runs[0].bold = True
        format_monochrome_table(tbl7, [1.8, 4.7])
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        
        # ---------------------------------------------------------
        # TABLA 8: GUÍA RÁPIDA DE TROUBLESHOOTING
        # ---------------------------------------------------------
        add_section_title(doc, "8. Guía de Resolución de Fallos Comunes (Troubleshooting)")
        
        t8_data = [("Problema / Síntoma Observado", "Causa Raíz Probable", "Acción Correctiva Inmediata")]
        for sint, acc in eq['troubleshooting']:
            t8_data.append((sint, "Desviación operativa o fatiga de componentes", acc))
            
        tbl8 = doc.add_table(rows=len(t8_data), cols=3)
        for r_i, (c0_txt, c1_txt, c2_txt) in enumerate(t8_data):
            tbl8.cell(r_i, 0).paragraphs[0].text = c0_txt
            tbl8.cell(r_i, 1).paragraphs[0].text = c1_txt
            tbl8.cell(r_i, 2).paragraphs[0].text = c2_txt
            if r_i > 0:
                tbl8.cell(r_i, 0).paragraphs[0].runs[0].bold = True
        format_monochrome_table(tbl8, [2.3, 1.8, 2.4])
        
        # Pie institucional
        p_foot = doc.add_paragraph()
        p_foot.paragraph_format.space_before = Pt(16)
        p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ft = p_foot.add_run(f"Proyecto de Investigación: Asistente Inteligente de Laboratorio — UTEQ 2026\nDocumento Técnico Individual N° {eq['num']} — Formato Monocromático de Alta Fidelidad")
        r_ft.font.name = "Calibri"
        r_ft.font.size = Pt(8.5)
        r_ft.font.color.rgb = COLOR_GRAY
        
        # Guardar
        doc.save(filepath)
        size = os.path.getsize(filepath)
        print(f"[{idx+1}/8] Generated: {filename} ({size} bytes)")
        
    print("All 8 individual documents generated successfully in knowledge_base/manuales_individuales/!")

if __name__ == "__main__":
    generate_individual_documents()
